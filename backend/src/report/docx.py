"""Editable DOCX twin of the PDF report (P5-2, design.md §9).

Renders the same ``ReportData`` snapshot with python-docx. The PDF remains
the authoritative rendering; this file exists so labs can edit/annotate
before official submission. Structure mirrors pdf.py section-for-section.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from .aggregate import ReportData, test_title
from .seal import qr_payload, qr_png_bytes

_ACCENT = RGBColor(0x1A, 0x3A, 0x5C)
_OBS_COLS = ("#", "Pos", "L", "I", "ΔL", "E", "Ec", "MPE", "MPE in e", "Verdict")


def _kv_table(doc: Document, pairs: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(pairs), cols=2)
    table.style = "Table Grid"
    for row, (label, value) in zip(table.rows, pairs):
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = str(value)
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def _obs_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(_OBS_COLS))
    table.style = "Table Grid"
    for col, title in zip(table.rows[0].cells, _OBS_COLS):
        col.text = title
        for r in col.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(8)
    for i, r in enumerate(rows, start=1):
        values = (r["seq"], r["position"], r["L"], r["I"], r["dL"], r["E"], r["Ec"], r["MPE"], r["MPE_e"], r["verdict_glyph"])
        for col, value in zip(table.rows[i].cells, values):
            col.text = value
            for p in col.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
                    run.font.name = "Courier New"


def render_docx(data: ReportData, *, verify_base_url: str, sha256: str, report_id: str) -> bytes:
    """Render the DOCX twin; returns complete .docx bytes."""
    doc = Document()
    doc.add_heading(data.lab["name"], level=0).runs[0].font.color.rgb = _ACCENT
    addr = doc.add_paragraph(data.lab["address"])
    addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Pattern Evaluation Report — Non-Automatic Weighing Instrument")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    clause = doc.add_paragraph("Compiled in accordance with OIML R 76-2 — " + str(data.overall["clause"]))
    clause.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in clause.runs:
        r.font.size = Pt(8)

    res = doc.add_paragraph()
    res.add_run("OVERALL RESULT: ").bold = True
    big = res.add_run(str(data.overall["result"]))
    big.bold = True
    big.font.size = Pt(20)

    doc.add_heading("1 · Instrument identity", level=1)
    _kv_table(doc, list(data.instrument.items()))
    doc.add_heading("2 · Laboratory & environmental conditions", level=1)
    _kv_table(doc, list(data.conditions.items()))
    doc.add_heading("3 · Personnel", level=1)
    _kv_table(
        doc,
        [
            ("Tested by", data.tested_by or "—"),
            ("Authorized Signatory", data.approved_by or "(pending officer sign-off)"),
        ],
    )

    doc.add_heading("4 · Verification seal", level=1)
    doc.add_picture(io.BytesIO(qr_png_bytes(qr_payload(verify_base_url, report_id, sha256))), width=Mm(34))
    doc.add_paragraph(f"Report ID: {report_id}")
    p = doc.add_paragraph("Content digest (SHA-256 of the report data): ")
    run = p.add_run(sha256)
    run.font.name = "Courier New"
    run.font.size = Pt(7)

    doc.add_heading("5 · Test observations", level=1)
    for i, (test_type, rows) in enumerate(data.observations.items(), start=1):
        fails = sum(1 for r in rows if r["verdict"] == "FAIL")
        doc.add_heading(f"5.{i} · {test_title(test_type)} — {len(rows)} reading(s), {fails} failed", level=2)
        _obs_table(doc, rows)

    if data.checklist:
        doc.add_heading("5.C · Checklist (R 76-2 sheet 17)", level=1)
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for c, t in zip(hdr, ("Clause", "Requirement", "Procedure", "Outcome", "Remarks")):
            c.text = t
        for item in data.checklist:
            cells = tbl.add_row().cells
            cells[0].text = item["clause"]
            cells[1].text = item["requirement"]
            cells[2].text = item["test_procedure"]
            cells[3].text = item["outcome"]
            cells[4].text = item["remarks"]
        prog = data.checklist_progress
        doc.add_paragraph(
            f"Checklist: {prog.get('passed', 0)} passed, {prog.get('failed', 0)} failed, "
            f"{prog.get('open', 0)} open, {prog.get('total', 0)} items."
        )

    checks = data.overall.get("checks") or []
    if checks:
        doc.add_heading("5.X · Criteria across readings", level=1)
        crit = doc.add_table(rows=1, cols=4)
        crit.style = "Table Grid"
        for c, t in zip(crit.rows[0].cells, ("Criterion", "Clause", "Result", "Detail")):
            c.text = t
        for check in checks:
            cells = crit.add_row().cells
            cells[0].text = check["title"]
            cells[1].text = check["clause"]
            cells[2].text = check["verdict"]
            cells[3].text = check["detail"]

    doc.add_heading("6 · Result summary", level=1)
    _kv_table(
        doc,
        [
            ("Readings evaluated", str(data.overall["total"])),
            ("Passed", str(data.overall["pass_count"])),
            ("Failed", str(data.overall["fail_count"])),
            ("Worst MPE utilization", str(data.overall["worst_utilization"])),
            ("Ambient drift (§3.9.2)", str(data.overall["drift_note"])),
            ("Evaluation basis", str(data.overall["clause"])),
        ],
    )
    doc.add_heading("Signatures", level=1)
    sig = doc.add_table(rows=1, cols=2)
    sig.style = "Table Grid"
    sig.rows[0].cells[0].text = "Tested by\n\n_________________________\n" + (data.tested_by or "—")
    sig.rows[0].cells[1].text = "Authorized Signatory\n\n_________________________\n" + (data.approved_by or "(pending sign-off)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


__all__ = ["render_docx"]
